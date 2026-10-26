"""CSV and Excel exports. Cells that start with = + - @ are prefixed with an
apostrophe so a spreadsheet never executes them as formulas. Digital codes
and buyer data other than the public nickname are never exported."""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ..database.enums import ORDER_STATUS_LABELS
from ..database.models import Order, Product
from ..shared.timeutil import fmt_dt
from . import analytics


def safe(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def to_csv(header: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(header)
    for row in rows:
        writer.writerow([safe("" if v is None else str(v)) for v in row])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def to_xlsx(sheets: list[tuple[str, list[str], list[list[Any]]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append([safe(v) for v in row])
        for idx, name in enumerate(header, start=1):
            width = max([len(str(name))] + [len(str(r[idx - 1])) for r in rows[:200] if idx - 1 < len(r) and r[idx - 1] is not None])
            ws.column_dimensions[get_column_letter(idx)].width = min(60, max(10, width + 2))
        ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


ORDER_HEADER = [
    "Номер",
    "Сделка PlayerOK",
    "Дата",
    "Статус",
    "Покупатель",
    "Товары",
    "Сумма, ₽",
    "Комиссия, ₽",
    "Себестоимость, ₽",
    "Чистая прибыль, ₽",
    "Тест",
    "Источник",
]


async def order_rows(session: AsyncSession, p: analytics.Period, include_test: bool) -> list[list[Any]]:
    stmt = select(Order).where(Order.placed_at >= p.start, Order.placed_at < p.end).order_by(Order.placed_at)
    if not include_test:
        stmt = stmt.where(Order.is_test.is_(False))
    rows = []
    for o in (await session.execute(stmt)).scalars():
        items = ", ".join(f"{i.product.name} × {i.quantity}" for i in o.items)
        rows.append(
            [
                o.number,
                o.external_id,
                fmt_dt(o.placed_at),
                ORDER_STATUS_LABELS[o.status].split(" ", 1)[1],
                o.buyer_name or "",
                items,
                o.total_amount,
                o.marketplace_fee,
                o.cost_total,
                o.net_profit,
                "да" if o.is_test else "нет",
                o.source,
            ]
        )
    return rows


async def analytics_export(session: AsyncSession, p: analytics.Period, fmt: str, include_test: bool = False) -> tuple[bytes, str, str]:
    m = await analytics.metrics(session, p, include_test=include_test)
    series = await analytics.timeseries(session, p, include_test=include_test)
    br = await analytics.breakdown(session, p, include_test=include_test)
    summary = [
        ["Период", m["period"]["label"]],
        ["Заказов", m["orders"]],
        ["Завершено сделок", m["completed"]],
        ["Отмен", m["cancelled"]],
        ["Возвратов", m["refunds"]],
        ["Споров", m["disputes"]],
        ["Отзывов", m["reviews"]],
        ["Выручка, ₽", m["revenue"]],
        ["Закупка, ₽", m["purchase_costs"]],
        ["Комиссии, ₽", m["fees"]],
        ["Чистая прибыль, ₽", m["net_profit"]],
        ["Маржинальность, %", m["margin_pct"]],
        ["Средний чек, ₽", m["avg_check"]],
        ["Среднее время до выдачи, мин", m["avg_processing_minutes"]],
        ["Сделок без себестоимости", m["orders_without_cost"]],
        ["Конверсия", m["conversion_note"]],
    ]
    days = [[d["day"], d["orders"], d["revenue"], d["net_profit"]] for d in series]
    products = [[x["name"], x["category"], x["quantity"], x["revenue"], x["profit_before_withdrawal"]] for x in br["products"]]
    suppliers = [[x["supplier"], x["orders"], x["succeeded"], x["failed"], x["success_rate"], x["spend_rub"]] for x in br["suppliers"]]
    orders = await order_rows(session, p, include_test)
    stamp = f"{p.first_day:%Y%m%d}-{p.last_day:%Y%m%d}"
    if fmt == "xlsx":
        data = to_xlsx(
            [
                ("Сводка", ["Показатель", "Значение"], summary),
                ("По дням", ["День", "Заказов", "Выручка, ₽", "Чистая прибыль, ₽"], days),
                ("Товары", ["Товар", "Категория", "Продано, шт", "Выручка, ₽", "Прибыль до вывода, ₽"], products),
                ("Поставщики", ["Поставщик", "Заказов", "Успешно", "Ошибок", "Успешность, %", "Расход, ₽"], suppliers),
                ("Заказы", ORDER_HEADER, orders),
            ]
        )
        return data, f"analytics-{stamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return to_csv(ORDER_HEADER, orders), f"orders-{stamp}.csv", "text/csv; charset=utf-8"


async def products_export(session: AsyncSession, fmt: str) -> tuple[bytes, str, str]:
    header = [
        "Артикул",
        "Название",
        "Категория",
        "Статус",
        "Себестоимость, ₽",
        "Мин. цена, ₽",
        "Рекомендуемая, ₽",
        "Текущая, ₽",
        "Прибыль, ₽",
        "Маржа, %",
        "Статус цены",
        "Рассчитано",
    ]
    rows = [
        [
            p.sku,
            p.name,
            p.category.name,
            p.status.value,
            p.full_cost,
            p.min_price,
            p.recommended_price,
            p.current_price,
            p.expected_profit,
            p.expected_margin_pct,
            p.pricing_status,
            fmt_dt(p.priced_at),
        ]
        for p in (await session.execute(select(Product).order_by(Product.id))).unique().scalars()
    ]
    if fmt == "xlsx":
        return to_xlsx([("Товары", header, rows)]), "products.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return to_csv(header, rows), "products.csv", "text/csv; charset=utf-8"
